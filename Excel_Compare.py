import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def generate_differences_for_all_cells(folder_path):
    # Find all Excel files in the folder
    all_files = [f for f in os.listdir(folder_path) if f.endswith('.xlsx')]

    # Ensure there are at least two Excel files in the folder
    if len(all_files) < 2:
        raise ValueError("The folder must contain at least two Excel files.")

    # Use the first two files in the folder
    file1 = os.path.join(folder_path, all_files[0])
    file2 = os.path.join(folder_path, all_files[1])

    print(f"Comparing files:\nFile 1: {file1}\nFile 2: {file2}")

   
    # Read the data into DataFrames
    df1 = pd.read_excel(file1, dtype=str)  # Read everything as strings for uniformity
    df2 = pd.read_excel(file2, dtype=str)

    df1 = df1.sort_values(by=list(df1.columns), ignore_index=True)
    df2 = df2.sort_values(by=list(df2.columns), ignore_index=True)


    # Ensure both DataFrames have the same column names
    all_columns = df1.columns.union(df2.columns)  # Combine all columns
    df1 = df1.reindex(columns=all_columns)  # Reindex File 1
    df2 = df2.reindex(columns=all_columns)  # Reindex File 2

    # Reset index to align rows
    df1.reset_index(drop=True, inplace=True)
    df2.reset_index(drop=True, inplace=True)

    # Expand DataFrames to have the same shape by filling missing rows with empty strings
    max_rows = max(len(df1), len(df2))
    df1 = df1.reindex(index=range(max_rows), fill_value="")
    df2 = df2.reindex(index=range(max_rows), fill_value="")

    # Create a DataFrame to store differences
    differences = pd.DataFrame(0, index=range(max_rows), columns=all_columns)

    # Compare cell by cell and calculate differences
    for col in all_columns:
        differences[col] = df1[col].where(df1[col] != df2[col], other="0")
        differences[col] = differences[col].mask(differences[col] != "0", pd.to_numeric(df1[col], errors='coerce') - pd.to_numeric(df2[col], errors='coerce'))

    # Check if there are no differences
    if differences.eq("0").all().all():
        print("No differences found between the two files.")
        return

    # Define output file name as differences_only.xlsx
    output_file_name = "differences_only.xlsx"
    output_file = os.path.join(folder_path, output_file_name)

    # Delete the file if it already exists
    if os.path.exists(output_file):
        os.remove(output_file)
        print(f"Existing file '{output_file}' deleted.")

    # Write the differences to a new Excel file
    differences.to_excel(output_file, index=False, sheet_name="Differences")

    # Highlight cells with differences in yellow
    wb = load_workbook(output_file)
    ws = wb["Differences"]
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value != "0":  # Highlight cells with differences
                cell.fill = yellow_fill

    wb.save(output_file)
    print(f"The file with highlighted differences is saved as '{output_file}'.")

# Example usage
folder_path = input("Enter the folder path containing the Excel files: ")
generate_differences_for_all_cells(folder_path)